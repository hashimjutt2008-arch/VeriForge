import csv
import zipfile
from collections.abc import Iterator
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook

MAX_CELL = 4096
MAX_COLUMNS = 256
MAX_EXPANDED = 128 * 1024 * 1024


@dataclass
class FileMetadata:
    filename: str
    columns: list[str]
    row_count: int
    suggested_column: int | None
    has_header: bool
    sample: list[list[str]]
    sheet: str | None = None

    def to_dict(self):
        return asdict(self)


def cell_text(value: object) -> str:
    text = "" if value is None else str(value)
    if len(text) > MAX_CELL:
        raise ValueError(f"A cell exceeds the {MAX_CELL}-character limit")
    return text


def text_stream(path: Path):
    with path.open("rb") as binary:
        bom = binary.read(2)
    return path.open(
        "r", encoding="utf-16" if bom in (b"\xff\xfe", b"\xfe\xff") else "utf-8-sig", newline=""
    )


def raw_rows(path: Path) -> Iterator[list[str]]:
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            with zipfile.ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED:
                    raise ValueError("Expanded spreadsheet exceeds the 128 MB safety limit")
                if len(archive.infolist()) > 10000:
                    raise ValueError("Spreadsheet contains too many archive entries")
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.worksheets[0]
                # Ignore misleading dimension hints in malformed workbooks.
                sheet.reset_dimensions()
                for row in sheet.iter_rows():
                    if len(row) > MAX_COLUMNS:
                        raise ValueError("Input exceeds 256 columns")
                    yield [cell_text(cell.value) for cell in row]
            finally:
                workbook.close()
        elif suffix in (".csv", ".txt"):
            with text_stream(path) as stream:
                if suffix == ".txt":
                    for line in stream:
                        yield [cell_text(line.rstrip("\r\n"))]
                else:
                    sample = stream.read(8192)
                    stream.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                    except csv.Error:
                        dialect = csv.excel
                    for row in csv.reader(stream, dialect, strict=True):
                        if len(row) > MAX_COLUMNS:
                            raise ValueError("Input exceeds 256 columns")
                        yield [cell_text(cell) for cell in row] or [""]
        else:
            raise ValueError("Supported formats are CSV, TXT and XLSX")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            "Could not read this file. Use a valid UTF-8/UTF-16 CSV/TXT or XLSX file."
        ) from exc


def inspect_file(
    path: Path, max_rows: int = 100000, has_header: bool | None = None
) -> FileMetadata:
    iterator = raw_rows(path)
    try:
        first = next(iterator, None)
        if first is None:
            raise ValueError("The file is empty")
        if has_header is None:
            has_header = path.suffix.lower() != ".txt" and not any("@" in value for value in first)
        if path.suffix.lower() == ".txt":
            has_header = False
        columns = []
        for i, value in enumerate(first):
            name = (value.strip()[:100] if has_header else "") or (
                "Email" if len(first) == 1 else f"Column {i + 1}"
            )
            if name in columns:
                name += f" ({i + 1})"
            columns.append(name)
        sample = [] if has_header else [first]
        count = 0 if has_header else 1
        for row in iterator:
            count += 1
            if count > max_rows:
                raise ValueError(f"Input exceeds the {max_rows:,}-row limit")
            if len(row) > len(columns):
                columns.extend(f"Column {i + 1}" for i in range(len(columns), len(row)))
            if len(sample) < 20:
                sample.append(row)
        if not count:
            raise ValueError("The file contains a header but no data rows")
        named = [
            i
            for i, name in enumerate(columns)
            if name.lower().replace("_", " ") in ("email", "email address", "e-mail")
        ]
        suggested = named[0] if len(named) == 1 else None
        if suggested is None:
            scores = [
                (i, sum(i < len(row) and "@" in row[i] for row in sample) / max(1, len(sample)))
                for i in range(len(columns))
            ]
            winners = [i for i, score in scores if score >= 0.8]
            if len(winners) == 1:
                suggested = winners[0]
        return FileMetadata(
            path.name,
            columns,
            count,
            suggested,
            has_header,
            sample,
            "First worksheet" if path.suffix.lower() == ".xlsx" else None,
        )
    finally:
        iterator.close()


def email_values(path: Path, metadata: FileMetadata, column: int) -> Iterator[str]:
    if not 0 <= column < len(metadata.columns):
        raise ValueError("Select an available email column")
    iterator = raw_rows(path)
    try:
        if metadata.has_header:
            next(iterator, None)
        for row in iterator:
            yield row[column] if column < len(row) else ""
    finally:
        iterator.close()
