import csv

import pytest
from openpyxl import Workbook, load_workbook

from engine.files import email_values, inspect_file
from engine.models import Options
from engine.report import Report


@pytest.mark.parametrize("kind", ["csv", "txt", "xlsx"])
def test_files_and_exports(tmp_path, kind):
    path = tmp_path / f"contacts.{kind}"
    values = ["first@corp.com", "second@corp.comt", "first@corp.com", "123@corp.com", "", "bad"]
    if kind == "xlsx":
        wb = Workbook()
        wb.active.append(["Name", "Email"])
        for i, value in enumerate(values):
            wb.active.append([f"Person {i}", value])
        wb.save(path)
    elif kind == "csv":
        with path.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.writer(stream)
            writer.writerow(["Name", "Email"])
            writer.writerows([f"Person {i}", value] for i, value in enumerate(values))
    else:
        path.write_text("\n".join(values) + "\n", encoding="utf-8")
    metadata = inspect_file(path)
    assert metadata.row_count == 6
    assert metadata.suggested_column == (0 if kind == "txt" else 1)
    report = Report(tmp_path / "report.sqlite")
    stats = report.process(email_values(path, metadata, metadata.suggested_column), Options())
    assert stats["total"] == 6 and stats["clean"] == 2
    expected = {"valid": 1, "corrected": 1, "removed": 3, "review": 1, "clean": 2, "full": 6}
    for export_type, count in expected.items():
        for format in ("csv", "xlsx"):
            target = tmp_path / f"{export_type}.{format}"
            assert report.export(target, export_type, format) == count
            if format == "csv":
                with target.open(encoding="utf-8-sig", newline="") as stream:
                    assert len(list(csv.reader(stream))) == count + 1
            else:
                wb = load_workbook(target, read_only=True)
                assert len(list(wb.active.values)) == count + 1
                wb.close()
    page = report.page(status="CORRECTED")
    assert page["total"] == 1
    assert page["items"][0]["row_number"] == 2
    assert page["items"][0]["original_email"] == values[1]
    assert report.page(search="second")["total"] == 1
    assert report.page(domain="corp.com", category="DUPLICATE")["total"] == 1


def test_no_header_bom_semicolon_and_missing_cells(tmp_path):
    path = tmp_path / "no-header.csv"
    path.write_text("Jane;jane@corp.com\nJoe;joe@corp.com\nBlank;\n", encoding="utf-8-sig")
    meta = inspect_file(path)
    assert meta.has_header is False
    assert list(email_values(path, meta, 1)) == ["jane@corp.com", "joe@corp.com", ""]
    assert meta.row_count == 3


def test_header_override_and_ambiguous_columns(tmp_path):
    path = tmp_path / "contacts.csv"
    path.write_text("one,two\na@corp.com,b@corp.com\n", encoding="utf-8")
    assert inspect_file(path).suggested_column is None
    assert inspect_file(path, has_header=False).row_count == 2


@pytest.mark.parametrize(
    "name,content",
    [
        ("empty.csv", b""),
        ("bad.xlsx", b"not-a-zip"),
        ("bad.csv", b"\xff\x80"),
        ("header.csv", b"email\n"),
        ("large.txt", b"a" * 4097),
        ("bad.exe", b"x"),
    ],
)
def test_bad_inputs(tmp_path, name, content):
    path = tmp_path / name
    path.write_bytes(content)
    with pytest.raises(ValueError):
        inspect_file(path)


def test_row_limit(tmp_path):
    path = tmp_path / "input.txt"
    path.write_text("a@x.com\nb@x.com\nc@x.com", encoding="utf-8")
    with pytest.raises(ValueError, match="row limit"):
        inspect_file(path, max_rows=2)


def test_xlsx_formulas_are_literal_strings(tmp_path):
    report = Report(tmp_path / "report.sqlite")
    report.process(["=1+1"], Options())
    out = tmp_path / "report.xlsx"
    report.export(out, "full", "xlsx")
    wb = load_workbook(out)
    assert wb.active["A2"].value == "=1+1"
    assert wb.active["A2"].data_type == "s"
    wb.close()


def test_streams_large_input_and_pages(tmp_path):
    report = Report(tmp_path / "report.sqlite")
    progress = []
    stats = report.process(
        (f"person{i}@gmail.com" for i in range(10000)),
        Options(),
        lambda count, counts: progress.append(count),
    )
    assert stats["clean"] == 10000
    assert progress[-1] == 10000
    assert len(report.page(offset=9950)["items"]) == 50
    assert report.page(search="' OR 1=1")["total"] == 0
